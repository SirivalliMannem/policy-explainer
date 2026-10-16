"""Text extraction for documents uploaded into the chat.

PDFs keep their pages; Word and text files are split into sections at their headings (or, without
headings, into blocks of about SECTION_CHARS), numbered from 1, so answers can cite "page 3" or
"section 2". Only the text is kept: the file itself is not stored.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass

MAX_BYTES = 10 * 1024 * 1024
MAX_PAGES = 500
SECTION_CHARS = 3000
MIN_TEXT_CHARS = 20  # less than this means no text layer (a scan) or an empty file

KINDS = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
    ".md": "text/markdown",
}


class ExtractionError(ValueError):
    """The file cannot be used; the message is shown to the employee."""


@dataclass
class Extracted:
    content_type: str
    unit: str  # page | section
    pages: list[str]


def _clean(text: str) -> str:
    text = (text or "").replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _blocks(paragraphs: list[tuple[bool, str]]) -> list[str]:
    """Sections from (is_heading, text) paragraphs: a heading starts a section; long ones are split."""
    sections: list[str] = []
    current = ""
    for is_heading, text in paragraphs:
        if not text.strip():
            continue
        if current and (is_heading or len(current) + len(text) > SECTION_CHARS):
            sections.append(current)
            current = ""
        current = f"{current}\n\n{text.strip()}".strip()
    if current:
        sections.append(current)
    return sections


def _pdf(data: bytes) -> list[str]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ExtractionError("The PDF is password-protected. Upload an unprotected copy.")
        if len(reader.pages) > MAX_PAGES:
            raise ExtractionError(f"The PDF has {len(reader.pages)} pages; the limit is {MAX_PAGES}.")
        return [_clean(page.extract_text() or "") for page in reader.pages]
    except ExtractionError:
        raise
    except (PdfReadError, ValueError, KeyError, OSError) as exc:
        raise ExtractionError("The PDF could not be read. It may be damaged.") from exc


def _docx(data: bytes) -> list[str]:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises several unrelated types for a bad file
        raise ExtractionError("The Word file could not be read. Save it again as .docx or as PDF.") from exc
    paragraphs = [
        ((p.style is not None and (p.style.name or "").lower().startswith(("heading", "title"))), p.text)
        for p in document.paragraphs
    ]
    for table in document.tables:  # schedules are often tables
        for row in table.rows:
            paragraphs.append((False, " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())))
    return [_clean(s) for s in _blocks(paragraphs)]


def _text(data: bytes) -> list[str]:
    text = data.decode("utf-8", errors="replace")
    paragraphs = [(bool(re.match(r"^\s*#{1,6}\s", block)), block) for block in re.split(r"\n\s*\n", text)]
    return [_clean(s) for s in _blocks(paragraphs)]


def extract(filename: str, data: bytes) -> Extracted:
    """The document's text per page or section, or ExtractionError with a reason to show."""
    name = (filename or "").lower()
    ext = next((e for e in KINDS if name.endswith(e)), None)
    if ext is None:
        raise ExtractionError("Only PDF, Word (.docx) and text (.txt, .md) files can be uploaded.")
    if not data:
        raise ExtractionError("The file is empty.")
    if len(data) > MAX_BYTES:
        raise ExtractionError(f"The file is {len(data) / 1024 / 1024:.1f} MB; the limit is {MAX_BYTES // 1024 // 1024} MB.")
    if ext == ".pdf":
        if not data.startswith(b"%PDF"):
            raise ExtractionError("The file is not a PDF.")
        pages, unit = _pdf(data), "page"
    elif ext == ".docx":
        if not data.startswith(b"PK"):
            raise ExtractionError("The file is not a Word .docx document.")
        pages, unit = _docx(data), "section"
    else:
        pages, unit = _text(data), "section"
    if sum(len(p) for p in pages) < MIN_TEXT_CHARS:
        raise ExtractionError(
            "No text could be read from this file. If it is a scanned document, upload a version with "
            "selectable text (scans without a text layer are not supported)."
        )
    if unit == "section":
        pages = [p for p in pages if p]  # sections are numbered as they read; PDF pages keep their numbers
    return Extracted(KINDS[ext], unit, pages)
