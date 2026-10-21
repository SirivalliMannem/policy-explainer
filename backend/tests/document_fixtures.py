"""Small documents built in code for the upload tests (no binary fixtures in the repository)."""

from __future__ import annotations

import io


def _escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def pdf(pages: list[list[str]]) -> bytes:
    """A minimal valid PDF: one page per list of text lines, with a real text layer."""
    objects: list[bytes] = []
    page_ids = [3 + 2 * i for i in range(len(pages))]
    font_id = 3 + 2 * len(pages)
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{pid} 0 R" for pid in page_ids)
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    for i, lines in enumerate(pages):
        content_id = page_ids[i] + 1
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode()
        )
        ops = ["BT", "/F1 11 Tf", "14 TL", "50 740 Td"]
        for line in lines:
            ops.append(f"({_escape(line)}) Tj T*")
        ops.append("ET")
        stream = "\n".join(ops).encode("latin-1")
        objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for n, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{n} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def docx(sections: list[tuple[str, list[str]]]) -> bytes:
    """A Word document: each section is a heading followed by paragraphs."""
    import docx as python_docx

    document = python_docx.Document()
    for heading, paragraphs in sections:
        document.add_heading(heading, level=1)
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


SAMPLE_POLICY_PAGES = [
    [
        "Lakeside Mutual Renters Policy - Declarations",
        "Policyholder: Jordan Avery. Policy period: January 1, 2026 to January 1, 2027.",
        "Personal property limit: $30,000. Deductible: $750 per claim.",
        "Loss of use: up to $6,000 while the rented home cannot be lived in.",
    ],
    [
        "Exclusions",
        "We do not pay for damage caused by flood or surface water.",
        "We do not pay for loss caused by mold unless it results from a covered water leak.",
        "Bicycles are covered up to $1,500 when stolen from a locked storage room.",
    ],
]
